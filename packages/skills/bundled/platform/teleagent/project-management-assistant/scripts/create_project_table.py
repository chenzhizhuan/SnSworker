"""
项目管理总表生成脚本
Usage: python create_project_table.py <output_path>
  如: python create_project_table.py "C:/path/to/项目管理工作区/项目管理总表.xlsx"
"""

import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation


def create(output_path: str):
    wb = Workbook()
    ws = wb.active
    ws.title = "项目总览"
    ws.sheet_properties.tabColor = "4472C4"

    headers = ["序号", "项目名称", "负责人", "阶段", "当前进度",
               "关键里程碑", "卡点问题", "资源需求", "状态", "风险等级", "最近更新", "备注"]
    col_widths = [6, 28, 10, 10, 12, 35, 30, 25, 10, 10, 14, 20]

    header_font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="4472C4")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin", color="B4C6E7"),
        right=Side(style="thin", color="B4C6E7"),
        top=Side(style="thin", color="B4C6E7"),
        bottom=Side(style="thin", color="B4C6E7")
    )

    for col_idx, (h, w) in enumerate(zip(headers, col_widths), 1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col_idx)].width = w

    ws.row_dimensions[1].height = 28

    # 预留10行
    body_font = Font(name="Arial", size=10)
    body_align = Alignment(vertical="center", wrap_text=True)
    alt_fill = PatternFill("solid", fgColor="D9E2F3")

    for row_idx in range(2, 12):
        ws.cell(row=row_idx, column=1, value=row_idx - 1)
        ws.cell(row=row_idx, column=9, value="🟢 正常")
        ws.cell(row=row_idx, column=10, value="低")
        for col_idx in range(1, 13):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = body_font
            cell.alignment = body_align
            cell.border = thin_border
            if row_idx % 2 == 0:
                cell.fill = alt_fill
        ws.row_dimensions[row_idx].height = 28

    # 冻结 & 筛选
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:L11"

    # 条件格式 - 状态列
    range_i = "I2:I11"
    ws.conditional_formatting.add(range_i, CellIsRule(
        operator="equal", formula=['"🟢 正常"'],
        fill=PatternFill("solid", fgColor="C6EFCE"),
        font=Font(color="006100", name="Arial", size=10)))
    ws.conditional_formatting.add(range_i, CellIsRule(
        operator="equal", formula=['"🟡 预警"'],
        fill=PatternFill("solid", fgColor="FFEB9C"),
        font=Font(color="9C5700", name="Arial", size=10)))
    ws.conditional_formatting.add(range_i, CellIsRule(
        operator="equal", formula=['"🔴 停滞"'],
        fill=PatternFill("solid", fgColor="FFC7CE"),
        font=Font(color="9C0006", name="Arial", size=10)))

    # 条件格式 - 风险列
    range_j = "J2:J11"
    ws.conditional_formatting.add(range_j, CellIsRule(
        operator="equal", formula=['"低"'],
        fill=PatternFill("solid", fgColor="C6EFCE"),
        font=Font(color="006100", name="Arial", size=10)))
    ws.conditional_formatting.add(range_j, CellIsRule(
        operator="equal", formula=['"中"'],
        fill=PatternFill("solid", fgColor="FFEB9C"),
        font=Font(color="9C5700", name="Arial", size=10)))
    ws.conditional_formatting.add(range_j, CellIsRule(
        operator="equal", formula=['"高"'],
        fill=PatternFill("solid", fgColor="FFC7CE"),
        font=Font(color="9C0006", name="Arial", size=10)))

    # 下拉菜单
    dv_stage = DataValidation(type="list",
        formula1='"需求阶段,设计阶段,开发阶段,测试阶段,验收阶段,已完结,暂停"',
        allow_blank=True)
    dv_stage.error = "请从下拉列表中选择阶段"
    dv_stage.errorTitle = "无效输入"
    dv_stage.prompt = "请选择项目阶段"
    dv_stage.promptTitle = "项目阶段"
    ws.add_data_validation(dv_stage)
    dv_stage.add("D2:D11")

    dv_status = DataValidation(type="list",
        formula1='"🟢 正常,🟡 预警,🔴 停滞"',
        allow_blank=True)
    dv_status.error = "请从下拉列表中选择状态"
    dv_status.errorTitle = "无效输入"
    ws.add_data_validation(dv_status)
    dv_status.add("I2:I11")

    dv_risk = DataValidation(type="list",
        formula1='"低,中,高"',
        allow_blank=True)
    ws.add_data_validation(dv_risk)
    dv_risk.add("J2:J11")

    # Sheet2: 使用说明
    ws2 = wb.create_sheet("使用说明")
    ws2.sheet_properties.tabColor = "70AD47"
    ws2.column_dimensions["A"].width = 80

    lines = [
        ("📋 项目管理总表 - 使用说明", Font(name="Arial", bold=True, size=14, color="4472C4")),
        ("", None),
        ("一、表格结构", Font(name="Arial", bold=True, size=11)),
        ("本表包含两个工作表：", None),
        ("  1.「项目总览」- 所有项目的统一看板", None),
        ("  2.「使用说明」- 本页", None),
        ("", None),
        ("二、各列说明", Font(name="Arial", bold=True, size=11)),
        ("  项目名称：填写项目全称", None),
        ("  负责人：项目经理或主要对接人", None),
        ("  阶段：从下拉菜单选择（需求/设计/开发/测试/验收/已完结/暂停）", None),
        ("  当前进度：填写百分比，手工更新", None),
        ("  关键里程碑：记录重要时间节点和交付物", None),
        ("  卡点问题：记录当前阻碍项目推进的问题", None),
        ("  资源需求：需要协调的人员、预算或其他资源", None),
        ("  状态：从下拉菜单选择（🟢正常/🟡预警/🔴停滞），自动着色", None),
        ("  风险等级：从下拉菜单选择（低/中/高），自动着色", None),
        ("  最近更新：记录最后一次更新日期", None),
        ("  备注：补充说明", None),
        ("", None),
        ("三、自动功能", Font(name="Arial", bold=True, size=11)),
        ("  - 状态和风险等级列有下拉菜单", None),
        ("  - 条件格式：选不同状态/风险会自动变色（绿/黄/红）", None),
        ("  - 冻结窗格：滚动时保持表头和项目名称可见", None),
        ("  - 自动筛选：可按阶段、状态、负责人等筛选", None),
        ("", None),
        ("四、建议", Font(name="Arial", bold=True, size=11)),
        ("  1. 建议每周更新一次进度和卡点", None),
        ("  2. 卡点问题写清楚「谁」需要「做什么」才能解决", None),
        ("  3. 状态为🔴停滞的项目需要优先关注", None),
        ("  4. 上传会议纪要或客户沟通记录，可让AI提取关键信息并更新本表", None),
    ]
    for row_idx, (text, font) in enumerate(lines, 1):
        cell = ws2.cell(row=row_idx, column=1, value=text)
        cell.font = font if font else Font(name="Arial", size=11)
        cell.alignment = Alignment(wrap_text=True)

    wb.save(output_path)
    print(f"Created: {output_path}")


if __name__ == "__main__":
    path = sys.argv[1] if len(sys.argv) > 1 else "项目管理总表.xlsx"
    create(path)
