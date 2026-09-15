#!/usr/bin/env python3
"""Build an editable deck on a native PowerPoint layout.

The script is intentionally content-neutral. Project-specific copy belongs in a
JSON spec, not in this reusable builder.

This is a starter builder, not a full implementation of the T1-T34 design
catalog. Supported slide types: cover, section, title_body, three_cards, kpi.
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any, Iterable

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


DEFAULT_COLORS = {
    "brand": "C00000",
    "brand_dark": "A8001A",
    "blue": "0070C0",
    "text": "262626",
    "muted": "888888",
    "surface": "FFFFFF",
    "surface_tint": "FFE4E4",
    "border": "E8B4B4",
    "positive": "16A34A",
}
DEFAULT_FONT = "Microsoft YaHei"
CONTENT_X = 0.45
CONTENT_Y = 1.15
CONTENT_W = 12.43
CONTENT_BOTTOM = 7.35


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build PPTX on a native template layout")
    parser.add_argument("--template", required=True, type=Path, help="Input .pptx template")
    parser.add_argument("--layout", default="2019-004", help="Slide layout name")
    parser.add_argument("--spec", required=True, type=Path, help="JSON deck spec")
    parser.add_argument("--output", required=True, type=Path, help="Output .pptx")
    parser.add_argument(
        "--keep-template-slides",
        action="store_true",
        help="Keep existing slides from the template; default removes them from the output copy",
    )
    parser.add_argument("--dry-run", action="store_true", help="Validate only; do not write output")
    return parser.parse_args()


def color(value: str) -> RGBColor:
    raw = value.strip().lstrip("#")
    if len(raw) != 6:
        raise ValueError(f"invalid RGB color: {value!r}")
    return RGBColor.from_string(raw.upper())


def load_spec(path: Path) -> dict[str, Any]:
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError as exc:
        raise ValueError(f"spec not found: {path}") from exc
    except json.JSONDecodeError as exc:
        raise ValueError(f"invalid JSON spec: {exc}") from exc
    if not isinstance(payload, dict):
        raise ValueError("spec root must be an object")
    slides = payload.get("slides")
    if not isinstance(slides, list) or not slides:
        raise ValueError("spec.slides must be a non-empty array")
    for index, slide in enumerate(slides, start=1):
        if not isinstance(slide, dict):
            raise ValueError(f"slide {index} must be an object")
        if slide.get("type") not in {"cover", "section", "title_body", "three_cards", "kpi"}:
            raise ValueError(f"slide {index} has unsupported type: {slide.get('type')!r}")
        if not slide.get("title"):
            raise ValueError(f"slide {index} requires title")
    return payload


def find_layout(prs: Presentation, name: str):
    matches = [layout for layout in prs.slide_layouts if layout.name == name]
    if not matches:
        available = ", ".join(layout.name for layout in prs.slide_layouts)
        raise ValueError(f"layout {name!r} not found; available: {available}")
    if len(matches) > 1:
        raise ValueError(f"layout name is not unique: {name!r}")
    return matches[0]


def remove_existing_slides(prs: Presentation) -> None:
    slide_ids = list(prs.slides._sldIdLst)  # python-pptx has no public delete API
    for slide_id in slide_ids:
        rel_id = slide_id.rId
        prs.part.drop_rel(rel_id)
        prs.slides._sldIdLst.remove(slide_id)


def add_textbox(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    text: str,
    *,
    size: float = 12,
    font: str = DEFAULT_FONT,
    font_color: RGBColor,
    bold: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    valign: MSO_ANCHOR = MSO_ANCHOR.MIDDLE,
    margin: float = 0,
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = valign
    frame.margin_left = Pt(margin)
    frame.margin_right = Pt(margin)
    frame.margin_top = Pt(margin)
    frame.margin_bottom = Pt(margin)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = str(text)
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = font_color
    return shape


def add_header_title(slide, prefix: str, title: str, colors: dict[str, RGBColor], font: str) -> None:
    shape = slide.shapes.add_textbox(Inches(0.80), Inches(0.24), Inches(11.0), Inches(0.46))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = Pt(0)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.LEFT
    for text, size, rgb in (
        (f"{prefix}：" if prefix else "", 18, colors["brand"]),
        (title, 16, colors["blue"]),
    ):
        if not text:
            continue
        run = paragraph.add_run()
        run.text = text
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = True
        run.font.color.rgb = rgb


def add_card(slide, x: float, y: float, w: float, h: float, colors: dict[str, RGBColor]):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = colors["surface"]
    shape.line.color.rgb = colors["border"]
    shape.line.width = Pt(1)
    shape.shadow.inherit = False
    return shape


def add_bullets(slide, items: Iterable[str], colors: dict[str, RGBColor], font: str) -> None:
    values = [str(item) for item in items]
    if not values:
        return
    box = slide.shapes.add_textbox(
        Inches(CONTENT_X + 0.2),
        Inches(CONTENT_Y + 0.75),
        Inches(CONTENT_W - 0.4),
        Inches(CONTENT_BOTTOM - CONTENT_Y - 1.0),
    )
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Pt(4)
    for index, item in enumerate(values):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = item
        paragraph.level = 0
        paragraph.space_after = Pt(9)
        paragraph.font.name = font
        paragraph.font.size = Pt(16)
        paragraph.font.color.rgb = colors["text"]


def build_cover(slide, spec: dict[str, Any], colors: dict[str, RGBColor], font: str) -> None:
    add_textbox(slide, 0.8, 1.7, 8.8, 1.0, spec["title"], size=34, font=font,
                font_color=colors["brand"], bold=True)
    if spec.get("subtitle"):
        add_textbox(slide, 0.8, 2.8, 9.5, 0.5, spec["subtitle"], size=18, font=font,
                    font_color=colors["blue"])
    if spec.get("kicker"):
        add_textbox(slide, 0.8, 3.45, 9.5, 0.4, spec["kicker"], size=13, font=font,
                    font_color=colors["muted"])


def build_section(slide, spec: dict[str, Any], colors: dict[str, RGBColor], font: str) -> None:
    add_header_title(slide, spec.get("prefix", "章节"), spec["title"], colors, font)
    add_textbox(slide, 0.8, 2.2, 11.7, 0.9, spec["title"], size=30, font=font,
                font_color=colors["brand"], bold=True, align=PP_ALIGN.CENTER)
    if spec.get("subtitle"):
        add_textbox(slide, 1.2, 3.3, 10.9, 0.5, spec["subtitle"], size=15, font=font,
                    font_color=colors["muted"], align=PP_ALIGN.CENTER)


def build_title_body(slide, spec: dict[str, Any], colors: dict[str, RGBColor], font: str) -> None:
    add_header_title(slide, spec.get("prefix", "内容"), spec["title"], colors, font)
    if spec.get("subtitle"):
        add_textbox(slide, CONTENT_X, CONTENT_Y, CONTENT_W, 0.45, spec["subtitle"], size=12,
                    font=font, font_color=colors["muted"])
    add_bullets(slide, spec.get("bullets", []), colors, font)


def build_three_cards(slide, spec: dict[str, Any], colors: dict[str, RGBColor], font: str) -> None:
    add_header_title(slide, spec.get("prefix", "要点"), spec["title"], colors, font)
    cards = spec.get("cards")
    if not isinstance(cards, list) or not 2 <= len(cards) <= 4:
        raise ValueError("three_cards requires 2-4 cards")
    gap = 0.25
    width = (CONTENT_W - gap * (len(cards) - 1)) / len(cards)
    for index, card in enumerate(cards):
        if not isinstance(card, dict) or not card.get("title"):
            raise ValueError("each card requires title")
        x = CONTENT_X + index * (width + gap)
        add_card(slide, x, CONTENT_Y + 0.45, width, 5.25, colors)
        add_textbox(slide, x + 0.15, CONTENT_Y + 0.65, width - 0.3, 0.45, card["title"],
                    size=15, font=font, font_color=colors["brand"], bold=True,
                    align=PP_ALIGN.CENTER)
        body = "\n".join(str(item) for item in card.get("items", []))
        add_textbox(slide, x + 0.18, CONTENT_Y + 1.25, width - 0.36, 3.9, body,
                    size=12, font=font, font_color=colors["text"], valign=MSO_ANCHOR.TOP,
                    margin=2)


def build_kpi(slide, spec: dict[str, Any], colors: dict[str, RGBColor], font: str) -> None:
    add_header_title(slide, spec.get("prefix", "数据"), spec["title"], colors, font)
    items = spec.get("items")
    if not isinstance(items, list) or not 2 <= len(items) <= 5:
        raise ValueError("kpi requires 2-5 items")
    gap = 0.22
    width = (CONTENT_W - gap * (len(items) - 1)) / len(items)
    for index, item in enumerate(items):
        if not isinstance(item, dict) or "value" not in item or "label" not in item:
            raise ValueError("each kpi item requires value and label")
        x = CONTENT_X + index * (width + gap)
        add_card(slide, x, 2.1, width, 2.5, colors)
        semantic = item.get("semantic", "brand")
        number_color = colors["positive"] if semantic == "positive" else colors["brand"]
        add_textbox(slide, x + 0.1, 2.45, width - 0.2, 0.8, item["value"], size=26,
                    font="Arial", font_color=number_color, bold=True, align=PP_ALIGN.CENTER)
        add_textbox(slide, x + 0.1, 3.35, width - 0.2, 0.45, item["label"], size=12,
                    font=font, font_color=colors["text"], align=PP_ALIGN.CENTER)
        if item.get("note"):
            add_textbox(slide, x + 0.1, 3.85, width - 0.2, 0.35, item["note"], size=10,
                        font=font, font_color=colors["muted"], align=PP_ALIGN.CENTER)


BUILDERS = {
    "cover": build_cover,
    "section": build_section,
    "title_body": build_title_body,
    "three_cards": build_three_cards,
    "kpi": build_kpi,
}


def main() -> int:
    args = parse_args()
    if not args.template.exists():
        print(f"ERROR: template not found: {args.template}", file=sys.stderr)
        return 2

    try:
        spec = load_spec(args.spec)
        prs = Presentation(args.template)
        layout = find_layout(prs, args.layout)
        theme_values = {**DEFAULT_COLORS, **spec.get("colors", {})}
        colors = {name: color(value) for name, value in theme_values.items()}
        font = spec.get("font", DEFAULT_FONT)

        if not args.keep_template_slides:
            remove_existing_slides(prs)

        for slide_spec in spec["slides"]:
            slide = prs.slides.add_slide(layout)
            BUILDERS[slide_spec["type"]](slide, slide_spec, colors, font)

        metadata = spec.get("metadata", {})
        if metadata.get("title"):
            prs.core_properties.title = str(metadata["title"])
        if metadata.get("author"):
            prs.core_properties.author = str(metadata["author"])

        if args.dry_run:
            print(f"VALID: {len(spec['slides'])} slides; layout={args.layout}")
            return 0

        args.output.parent.mkdir(parents=True, exist_ok=True)
        prs.save(args.output)
        print(f"BUILT: {args.output} ({len(spec['slides'])} slides; layout={args.layout})")
        return 0
    except (OSError, ValueError, KeyError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
