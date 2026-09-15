#!/usr/bin/env python3
"""Analyze a PPTX template into a deterministic Layout Map.

Optionally exports real slide screenshots through Microsoft PowerPoint on
Windows. Font fields reflect explicit runs only; missing values mean the style
is inherited from the layout/master/theme and must be verified visually.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import sys
from pathlib import Path
from typing import Any


EMU_PER_INCH = 914400


def emu_to_inches(value: int | None) -> float | None:
    return round(value / EMU_PER_INCH, 4) if value is not None else None


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def ratio_label(width: float, height: float) -> tuple[str, float]:
    value = width / height
    if abs(value - (16 / 9)) <= 0.02:
        return "16:9", round(value, 6)
    if abs(value - (4 / 3)) <= 0.02:
        return "4:3", round(value, 6)
    return "custom", round(value, 6)


def explicit_font_info(run: Any) -> dict[str, Any]:
    font = run.font
    color = None
    try:
        if font.color and font.color.rgb:
            color = f"#{font.color.rgb}"
    except (AttributeError, TypeError, ValueError):
        color = None
    return {
        "source": "explicit-run",
        "name": font.name,
        "size": round(font.size.pt, 1) if font.size else None,
        "bold": font.bold,
        "italic": font.italic,
        "underline": font.underline,
        "color": color,
    }


def placeholder_info(placeholder: Any) -> dict[str, Any]:
    info: dict[str, Any] = {
        "idx": placeholder.placeholder_format.idx,
        "name": placeholder.name,
        "type": str(placeholder.placeholder_format.type),
        "x": emu_to_inches(placeholder.left),
        "y": emu_to_inches(placeholder.top),
        "w": emu_to_inches(placeholder.width),
        "h": emu_to_inches(placeholder.height),
    }
    if not placeholder.has_text_frame:
        return info

    frame = placeholder.text_frame
    info["text_content"] = frame.text or None
    if frame.paragraphs:
        info["alignment"] = str(frame.paragraphs[0].alignment) if frame.paragraphs[0].alignment else None
    for paragraph in frame.paragraphs:
        for run in paragraph.runs:
            if run.text.strip():
                info["font"] = explicit_font_info(run)
                return info
    info["font"] = {
        "source": "inherited",
        "name": None,
        "size": None,
        "bold": None,
        "italic": None,
        "underline": None,
        "color": None,
    }
    return info


def fill_info(element: Any) -> dict[str, Any]:
    fill = element.fill
    if fill.type is None:
        return {"type": "inherit-or-none"}
    fill_type = str(fill.type)
    if "SOLID" in fill_type:
        color = None
        try:
            if fill.fore_color.rgb:
                color = f"#{fill.fore_color.rgb}"
        except (AttributeError, TypeError, ValueError):
            color = None
        return {"type": "solid", "color": color}
    if "GRADIENT" in fill_type:
        stops = []
        try:
            for stop in fill.gradient_stops:
                stop_color = f"#{stop.color.rgb}" if stop.color and stop.color.rgb else None
                stops.append({"position": round(stop.position, 3), "color": stop_color})
        except (AttributeError, TypeError, ValueError):
            pass
        return {"type": "gradient", "stops": stops}
    if "PATTERN" in fill_type:
        return {"type": "pattern"}
    if "PICTURE" in fill_type:
        return {"type": "picture"}
    return {"type": fill_type}


def shape_info(shape: Any, index: int) -> dict[str, Any]:
    info: dict[str, Any] = {
        "index": index,
        "shape_type": str(shape.shape_type),
        "name": shape.name,
        "x": emu_to_inches(shape.left),
        "y": emu_to_inches(shape.top),
        "w": emu_to_inches(shape.width),
        "h": emu_to_inches(shape.height),
    }
    try:
        info["fill"] = fill_info(shape)
    except (AttributeError, TypeError, ValueError):
        info["fill"] = {"type": "unknown"}
    try:
        if shape.line and shape.line.color and shape.line.color.rgb:
            info["border"] = {
                "color": f"#{shape.line.color.rgb}",
                "width_pt": round(shape.line.width.pt, 1) if shape.line.width else None,
            }
    except (AttributeError, TypeError, ValueError):
        pass
    return info


def analyze_template(input_path: Path, sample_slides: int) -> dict[str, Any]:
    from pptx import Presentation

    presentation = Presentation(input_path)
    width = round(presentation.slide_width / EMU_PER_INCH, 3)
    height = round(presentation.slide_height / EMU_PER_INCH, 3)
    aspect_ratio, aspect_ratio_value = ratio_label(width, height)
    layout_map: dict[str, Any] = {
        "version": 2,
        "template_name": input_path.stem,
        "source_file": input_path.name,
        "source_sha256": sha256(input_path),
        "slide_width": width,
        "slide_height": height,
        "aspect_ratio": aspect_ratio,
        "aspect_ratio_value": aspect_ratio_value,
        "font_extraction_note": "Only explicit run fonts are reported; null values inherit from layout/master/theme.",
        "layouts": [],
        "actual_slides_sample": [],
    }

    for index, layout in enumerate(presentation.slide_layouts):
        layout_item: dict[str, Any] = {
            "layout_index": index,
            "layout_name": layout.name,
            "placeholders": [placeholder_info(ph) for ph in layout.placeholders],
            "background": {"type": "inherit"},
            "decorations": [],
        }
        for shape_index, shape in enumerate(shape for shape in layout.shapes if not shape.is_placeholder):
            layout_item["decorations"].append(shape_info(shape, shape_index))
        try:
            layout_item["background"] = fill_info(layout.background)
        except (AttributeError, TypeError, ValueError):
            pass
        layout_map["layouts"].append(layout_item)

    for index, slide in enumerate(presentation.slides):
        if sample_slides > 0 and index >= sample_slides:
            break
        item = {
            "slide_index": index + 1,
            "layout_name": slide.slide_layout.name if slide.slide_layout else None,
            "texts": [],
        }
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                item["texts"].append(
                    {
                        "name": shape.name,
                        "text": shape.text_frame.text[:200],
                        "is_placeholder": shape.is_placeholder,
                    }
                )
        layout_map["actual_slides_sample"].append(item)
    return layout_map


def export_screenshots(input_path: Path, output_dir: Path, width_px: int = 1600) -> list[str]:
    if os.name != "nt":
        raise RuntimeError("--screenshots requires Microsoft PowerPoint on Windows")
    try:
        import win32com.client  # type: ignore[import-not-found]
    except ImportError as exc:
        raise RuntimeError("--screenshots requires pywin32 and Microsoft PowerPoint") from exc

    output_dir.mkdir(parents=True, exist_ok=True)
    app = None
    presentation = None
    exported: list[str] = []
    try:
        app = win32com.client.DispatchEx("PowerPoint.Application")
        app.Visible = 1
        presentation = app.Presentations.Open(str(input_path.resolve()), True, False, False)
        ratio = float(presentation.PageSetup.SlideHeight) / float(presentation.PageSetup.SlideWidth)
        height_px = max(1, round(width_px * ratio))
        for index in range(1, presentation.Slides.Count + 1):
            path = output_dir / f"slide-{index:02d}.png"
            presentation.Slides.Item(index).Export(str(path.resolve()), "PNG", width_px, height_px)
            exported.append(str(path.resolve()))
    except Exception as exc:  # COM surfaces implementation-specific exception classes.
        raise RuntimeError(f"PowerPoint screenshot export failed: {exc}") from exc
    finally:
        if presentation is not None:
            presentation.Close()
        if app is not None:
            app.Quit()
    return exported


def main() -> int:
    parser = argparse.ArgumentParser(description="Extract a PPTX Layout Map")
    parser.add_argument("--input", required=True, type=Path, help="Input .pptx template")
    parser.add_argument("--output", required=True, type=Path, help="Output Layout Map JSON")
    parser.add_argument("--screenshots", type=Path, help="Optional directory for real PowerPoint slide PNGs")
    parser.add_argument(
        "--sample-slides",
        type=int,
        default=8,
        help="Number of actual slides to sample; 0 means all (default: 8)",
    )
    args = parser.parse_args()

    if not args.input.is_file():
        print(f"ERROR: file not found: {args.input}", file=sys.stderr)
        return 2
    if args.sample_slides < 0:
        print("ERROR: --sample-slides must be >= 0", file=sys.stderr)
        return 2
    try:
        layout_map = analyze_template(args.input, args.sample_slides)
        if args.screenshots:
            layout_map["screenshots"] = export_screenshots(args.input, args.screenshots)
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(json.dumps(layout_map, ensure_ascii=False, indent=2), encoding="utf-8")
    except ImportError:
        print("ERROR: python-pptx is required", file=sys.stderr)
        return 2
    except (OSError, RuntimeError, ValueError) as exc:
        print(f"ERROR: template analysis failed: {exc}", file=sys.stderr)
        return 1

    print(f"TEMPLATE_ANALYZED {layout_map['template_name']}")
    print(
        f"canvas={layout_map['slide_width']}x{layout_map['slide_height']}in "
        f"ratio={layout_map['aspect_ratio']} layouts={len(layout_map['layouts'])}"
    )
    print(f"output={args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
