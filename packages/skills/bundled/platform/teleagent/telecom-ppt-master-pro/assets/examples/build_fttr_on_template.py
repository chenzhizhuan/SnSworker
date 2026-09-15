# -*- coding: utf-8 -*-
"""
基于电信5G原生模板1.0.pptx 构建FTTR培训12页PPT
- 保留原版master/layout不动，每页用 '2019-004' layout 继承原生页眉
  （红色方块叠块 + 5G logo图片 + 红色横线）
- 页眉左侧补标题文字（前缀红：后缀蓝），正文区填内容
- 不复刻页眉，直接用底版原生元素
"""
import os, copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu, Length
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
SKILL_DIR = os.path.dirname(SCRIPT_DIR)
SRC = os.path.join(SKILL_DIR, "assets", "电信5G原生模板1.0.pptx")
OUT = os.path.abspath("FTTR全屋光网产品卖点培训_底版版.pptx")

# 电信色
RED   = RGBColor(0xD4, 0x12, 0x24)   # 模板原红 D41224
RED2  = RGBColor(0xC0, 0x00, 0x00)
BLUE  = RGBColor(0x00, 0x70, 0xC0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x33, 0x33, 0x33)
MUTED = RGBColor(0x88, 0x88, 0x88)
ACCENT= RGBColor(0xF2, 0xC3, 0x00)   # 金黄
ICELT = RGBColor(0xFD, 0xEC, 0xEC)   # 浅红
ICEMID= RGBColor(0xF7, 0xDC, 0xDC)
ICEBL = RGBColor(0xEA, 0xF1, 0xFB)   # 浅蓝
BORDER= RGBColor(0xE5, 0xE5, 0xE5)
GREEN = RGBColor(0x16, 0xA3, 0x4A)
FONT  = "微软雅黑"

# 正文安全区（页眉横线 y=0.83，留呼吸）
SX, SW = Inches(0.45), Inches(12.43)
SY     = Inches(1.15)

def add_textbox(s, x, y, w, h, runs, align="left", valign="middle", line_spacing=None):
    """runs: list of (text, opts) ; opts: dict size/bold/color/font"""
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(0); tf.margin_right = Pt(0)
    tf.margin_top = Pt(0); tf.margin_bottom = Pt(0)
    if valign == "middle": tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    elif valign == "top": tf.vertical_anchor = MSO_ANCHOR.TOP
    if align == "center": p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    else: p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
    if line_spacing: p.line_spacing = line_spacing
    first = True
    for text, opts in runs:
        if first:
            p.text = text if text else ""
            first = False
        else:
            p.add_run().text = text
    # apply format per run
    for i, (text, opts) in enumerate(runs):
        run = p.runs[i]
        run.font.name = opts.get("font", FONT)
        run.font.size = Pt(opts.get("size", 12))
        if opts.get("bold"): run.font.bold = True
        col = opts.get("color", BLACK)
        if col is not None: run.font.color.rgb = col
    return tb

def add_line_runs(s, x, y, w, lines, align="left", valign="top"):
    """多行，每行一个paragraph; lines=[(txt,opts)] each line independent"""
    tb = s.shapes.add_textbox(x, y, w, Inches(0.01))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left=Pt(0);tf.margin_right=Pt(0);tf.margin_top=Pt(0);tf.margin_bottom=Pt(0)
    if valign=="middle": tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    for i,(text,opts) in enumerate(lines):
        p = tf.paragraphs[0] if i==0 else tf.add_paragraph()
        if align=="center": p.alignment=PP_ALIGN.CENTER
        r=p.add_run(); r.text=text
        r.font.name=opts.get("font",FONT)
        r.font.size=Pt(opts.get("size",12))
        if opts.get("bold"): r.font.bold=True
        col=opts.get("color",BLACK)
        if col is not None: r.font.color.rgb=col
    return tb

def add_rect(s, x, y, w, h, fill=None, line=None, line_w=None, shape=MSO_SHAPE.RECTANGLE, radius=None):
    sp = s.shapes.add_shape(shape, x, y, w, h)
    if fill is not None:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    else:
        sp.fill.background()
    if line is not None:
        sp.line.color.rgb = line
        if line_w is None:
            sp.line.width = Pt(1)
        elif isinstance(line_w, Length):
            sp.line.width = line_w
        elif isinstance(line_w, (int, float)):
            sp.line.width = Pt(line_w)
        else:
            sp.line.width = line_w
    else:
        sp.line.fill.background()
    sp.shadow.inherit = False
    return sp

def add_oval(s, x, y, w, h, fill=None, line=None, line_w=None):
    return add_rect(s, x, y, w, h, fill, line, line_w, MSO_SHAPE.OVAL)

def add_header_title(s, prefix, suffix):
    """页眉标题文字（在原生红色方块右侧、横线上方区域 y~0.26-0.64, x~0.8）"""
    add_textbox(s, Inches(0.80), Inches(0.24), Inches(11.0), Inches(0.46),
                [(prefix+"：", {"size":18,"bold":True,"color":RED2}),
                 (suffix, {"size":16,"bold":True,"color":BLUE})],
                align="left", valign="middle")

def draw_card(s, x, y, w, h, line=RED2, line_w=1.0, fill=WHITE, radius=None):
    add_rect(s, x, y, w, h, fill=fill, line=line, line_w=Pt(line_w), shape=MSO_SHAPE.ROUNDED_RECTANGLE)

def make_shadow_left_xml(sp):
    """给shape加阴影的简单oxml"""
    spPr = sp._element.spPr
    # remove existing effectLst
    for tag in ('a:effectLst',):
        e = spPr.find(qn(tag))
        if e is not None: spPr.remove(e)
    eff = spPr.makeelement(qn('a:effectLst'), {})
    sh = eff.makeelement(qn('a:outerShdw'), {'blurRad':'40000','dist':'20000','dir':'5400000','rotWithShape':'0'})
    clr = sh.makeelement(qn('a:srgbClr'), {'val':'BFBFBF'})
    alpha = clr.makeelement(qn('a:alpha'), {'val':'40000'})
    clr.append(alpha); sh.append(clr); eff.append(sh); spPr.append(eff)

# ============================ 构建 ============================
prs = Presentation(SRC)
# 找到 2019-004 layout
TARGET_LAYOUT_NAME = "2019-004"
layout = None
for lo in prs.slide_masters[0].slide_layouts:
    if lo.name == TARGET_LAYOUT_NAME:
        layout = lo; break
if layout is None:
    raise RuntimeError("找不到 2019-004 layout")

# 移除原 slide1 (xxxx占位)
xml_slides = prs.slides._sldIdLst
slides = list(xml_slides)
xml_slides.remove(slides[0])

def new_slide():
    return prs.slides.add_slide(layout)

print("Building 12 slides on 2019-004 layout (native header preserved)...")

# ---------- P1 封面 ----------
s = new_slide()
# 封面用全幅，不画页眉标题文字（保留原生红方块+5G logo+横线作为顶部装饰）
# 主标题
add_textbox(s, Inches(0.8), Inches(1.7), Inches(8.5), Inches(1.1),
            [("FTTR 全屋光网", {"size":46,"bold":True,"color":RED2})], valign="middle")
add_textbox(s, Inches(0.8), Inches(2.8), Inches(8.5), Inches(0.45),
            [("Fiber to the Room", {"size":21,"bold":True,"color":BLUE})])
add_textbox(s, Inches(0.8), Inches(3.3), Inches(9), Inches(0.45),
            [("产品卖点培训  ·  千兆到房间  无盲区覆盖", {"size":16,"color":BLACK})])
# 三个chips
chips=["千兆到房间","全屋无盲区","无缝漫游"]
for i,c in enumerate(chips):
    cx=Inches(0.8+i*2.3)
    add_rect(s, cx, Inches(3.95), Inches(2.1), Inches(0.42), fill=ICELT, line=RED2, line_w=Pt(1), shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_textbox(s, cx, Inches(3.95), Inches(2.1), Inches(0.42), [(c,{"size":12,"bold":True,"color":RED2})], align="center", valign="middle")
# 右侧WiFi同心圆
dCx=10.6; dCy=2.9
add_rect(s, Inches(dCx-1.6), Inches(dCy-1.6), Inches(3.2), Inches(3.2), fill=ICELT, line=ACCENT, line_w=Pt(2), shape=MSO_SHAPE.OVAL)
add_oval(s, Inches(dCx-1.05), Inches(dCy-1.05), Inches(2.1), Inches(2.1), fill=WHITE, line=RED2, line_w=Pt(2))
add_oval(s, Inches(dCx-0.55), Inches(dCy-0.55), Inches(1.1), Inches(1.1), fill=ICELT, line=RED2, line_w=Pt(2))
add_oval(s, Inches(dCx-0.2), Inches(dCy-0.2), Inches(0.4), Inches(0.4), fill=RED2)
add_textbox(s, Inches(dCx-1.6), Inches(dCy+1.65), Inches(3.2), Inches(0.35), [("全屋覆盖",{"size":12,"bold":True,"color":MUTED})], align="center")
# 下半4指标条
metrics=[("1000M","到房速率"),("1拖N","全屋房间覆盖"),("0死角","信号盲区"),("无缝漫游","切换不掉线")]
mY=5.3; mGap=0.3; mW=(12.43-3*mGap)/4
for i,(n,l) in enumerate(metrics):
    mx=0.45+i*(mW+mGap)
    draw_card(s, Inches(mx), Inches(mY), Inches(mW), Inches(1.25), line=BORDER, line_w=1)
    add_textbox(s, Inches(mx), Inches(mY+0.2), Inches(mW), Inches(0.5), [(n,{"size":18,"bold":True,"color":RED2})], align="center", valign="middle")
    add_textbox(s, Inches(mx), Inches(mY+0.75), Inches(mW), Inches(0.35), [(l,{"size":11,"color":MUTED})], align="center", valign="middle")

# ---------- 通用页眉标题辅助 ----------
def hdr(s, prefix, suffix): add_header_title(s, prefix, suffix)

# ---------- P2 目录 ----------
s = new_slide(); hdr(s,"培训导航","目录 Contents")
add_textbox(s, SX, SY, SW, Inches(0.4), [("一图掌握FTTR培训四大模块",{"size":12,"color":MUTED})])
toc=[("01","认识FTTR","概念·架构·跃迁",RED2),("02","核心卖点","覆盖·速率·体验",BLUE),
     ("03","应用场景","5大典型场景",RED2),("04","销售打法","客群·话术·套餐",BLUE)]
tY=1.85; tH=2.4; twoW=(12.43-0.4)/2
for i,(no,title,desc,col) in enumerate(toc):
    col_i=i%2; row=i//2
    cx=0.45+col_i*(twoW+0.4); cy=tY+row*(tH+0.3)
    draw_card(s, Inches(cx), Inches(cy), Inches(twoW), Inches(tH), line=BORDER, line_w=1)
    # 编号块
    add_rect(s, Inches(cx), Inches(cy), Inches(1.4), Inches(tH), fill=col)
    add_textbox(s, Inches(cx), Inches(cy+0.3), Inches(1.4), Inches(1.0), [(no,{"size":40,"bold":True,"color":WHITE,"font":FONT})], align="center", valign="middle")
    add_textbox(s, Inches(cx+1.6), Inches(cy+0.5), Inches(twoW-1.8), Inches(0.55), [(title,{"size":20,"bold":True,"color":col})], valign="middle")
    add_textbox(s, Inches(cx+1.6), Inches(cy+1.15), Inches(twoW-1.8), Inches(0.5), [(desc,{"size":13,"color":MUTED})], valign="middle")
    add_rect(s, Inches(cx+1.6), Inches(cy+tH-0.5), Inches(0.8), Inches(0.04), fill=col)

# ---------- P3 什么是FTTR ----------
s = new_slide(); hdr(s,"认识FTTR","什么是FTTR")
# 定义条
add_rect(s, SX, SY, SW, Inches(0.55), fill=RED2, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
add_textbox(s, Inches(0.65), SY, Inches(12.0), Inches(0.55),
            [("FTTR ",{"size":14,"bold":True,"color":WHITE}),
             ("(Fiber to the Room) ",{"size":11,"color":WHITE}),
             ("光纤到房间——在光纤到户基础上，将光纤直接延伸至每个房间，实现全屋双千兆覆盖",{"size":12,"bold":True,"color":WHITE})], valign="middle")
# 左：核心价值
leftW=5.6
draw_card(s, SX, Inches(2.05), Inches(leftW), Inches(5.0))
add_rect(s, SX, Inches(2.05), Inches(leftW), Inches(0.5), fill=ICEMID)
add_textbox(s, Inches(0.7), Inches(2.05), Inches(leftW-0.5), Inches(0.5), [("核心价值",{"size":14,"bold":True,"color":RED2})], valign="middle")
vals=[("解决最后十米","光纤通到每个房间，突破室内带宽瓶颈"),("双千兆覆盖","有线无线双千兆，网速不再打折"),("全屋无盲区","WiFi信号无死角，告别穿墙衰减")]
for i,(t,d) in enumerate(vals):
    vy=2.75+i*1.35
    add_rect(s, Inches(0.7), Inches(vy), Inches(0.08), Inches(1.1), fill=RED2)
    add_textbox(s, Inches(0.95), Inches(vy), Inches(leftW-1.3), Inches(0.5), [(t,{"size":14,"bold":True,"color":BLACK})], valign="middle")
    add_textbox(s, Inches(0.95), Inches(vy+0.5), Inches(leftW-1.3), Inches(0.55), [(d,{"size":11,"color":MUTED})], valign="top")
# 右：FTTH vs FTTR
rxR=0.45+leftW+0.35; rW=12.43-leftW-0.35
draw_card(s, Inches(rxR), Inches(2.05), Inches(rW), Inches(5.0))
add_rect(s, Inches(rxR), Inches(2.05), Inches(rW), Inches(0.5), fill=ICEMID)
add_textbox(s, Inches(rxR+0.25), Inches(2.05), Inches(rW-0.5), Inches(0.5), [("FTTH  vs  FTTR",{"size":14,"bold":True,"color":RED2})], valign="middle")
# 对比子卡
cmpY=2.65; cmpH=4.3; subW=(rW-0.3)/2
# FTTH 子卡（左，灰）
ftx=rxR
add_rect(s, Inches(ftx), Inches(cmpY), Inches(subW), Inches(cmpH), fill=RGBColor(0xF5,0xF5,0xF5), line=RGBColor(0xD4,0xD4,0xD4), line_w=Pt(1), shape=MSO_SHAPE.ROUNDED_RECTANGLE)
add_textbox(s, Inches(ftx), Inches(cmpY+0.2), Inches(subW), Inches(0.45), [("FTTH",{"size":18,"bold":True,"color":MUTED,"font":FONT})], align="center", valign="middle")
add_textbox(s, Inches(ftx), Inches(cmpY+0.65), Inches(subW), Inches(0.35), [("光纤到户",{"size":12,"color":MUTED})], align="center", valign="middle")
ftth_pts=[("光纤到门口","室内靠网线/无线接力"),("穿墙衰减","远端房间速率打折"),("多设备争抢","并发能力有限")]
for i,(t,d) in enumerate(ftth_pts):
    py=cmpY+1.15+i*1.0
    add_oval(s, Inches(ftx+0.25), Inches(py+0.08), Inches(0.12), Inches(0.12), fill=MUTED)
    add_textbox(s, Inches(ftx+0.45), Inches(py), Inches(subW-0.6), Inches(0.4), [(t,{"size":12,"bold":True,"color":RGBColor(0x66,0x66,0x66)})], valign="middle")
    add_textbox(s, Inches(ftx+0.45), Inches(py+0.38), Inches(subW-0.6), Inches(0.4), [(d,{"size":10,"color":MUTED})], valign="middle")
# FTTR 子卡（右，红边）
frx=rxR+subW+0.3
add_rect(s, Inches(frx), Inches(cmpY), Inches(subW), Inches(cmpH), fill=WHITE, line=RED2, line_w=1.5, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
add_rect(s, Inches(frx), Inches(cmpY), Inches(subW), Inches(0.4), fill=RED2, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
add_textbox(s, Inches(frx), Inches(cmpY+0.05), Inches(subW), Inches(0.35), [("FTTR",{"size":16,"bold":True,"color":WHITE,"font":FONT})], align="center", valign="middle")
add_textbox(s, Inches(frx), Inches(cmpY+0.5), Inches(subW), Inches(0.35), [("光纤到房间",{"size":12,"bold":True,"color":RED2})], align="center", valign="middle")
fttr_pts=[("光纤到每个房间","室内全光千兆不打折"),("无穿墙衰减","全屋信号满格无死角"),("多设备并发","WiFi6低时延稳连接")]
for i,(t,d) in enumerate(fttr_pts):
    py=cmpY+1.0+i*1.05
    add_oval(s, Inches(frx+0.25), Inches(py+0.08), Inches(0.12), Inches(0.12), fill=RED2)
    add_textbox(s, Inches(frx+0.45), Inches(py), Inches(subW-0.6), Inches(0.4), [(t,{"size":12,"bold":True,"color":RED2})], valign="middle")
    add_textbox(s, Inches(frx+0.45), Inches(py+0.38), Inches(subW-0.6), Inches(0.4), [(d,{"size":10,"color":BLACK})], valign="middle")

# ---------- P4 组网架构 ----------
s = new_slide(); hdr(s,"认识FTTR","FTTR组网架构")
add_textbox(s, SX, SY, SW, Inches(0.4), [("1拖N组网：1台主光猫 + N台从光猫，隐形光纤串联，全屋双千兆",{"size":12,"color":MUTED})])
# 主光猫
mainX=0.9; mainY=1.95; mainW=2.2; mainH=1.3
draw_card(s, Inches(mainX), Inches(mainY), Inches(mainW), Inches(mainH), line=RED2, line_w=1.5)
add_rect(s, Inches(mainX), Inches(mainY), Inches(mainW), Inches(0.35), fill=RED2)
add_textbox(s, Inches(mainX), Inches(mainY+0.05), Inches(mainW), Inches(0.3), [("主光猫 WiFi6",{"size":11,"bold":True,"color":WHITE})], align="center", valign="middle")
add_textbox(s, Inches(mainX), Inches(mainY+0.45), Inches(mainW), Inches(0.4), [("千兆网关",{"size":14,"bold":True,"color":BLACK})], align="center", valign="middle")
add_textbox(s, Inches(mainX), Inches(mainY+0.9), Inches(mainW), Inches(0.3), [("统一接入调度",{"size":10,"color":MUTED})], align="center", valign="middle")
# 主光纤横线
fiberX=mainX+mainW; fiberEndX=12.4; fiberY=mainY+mainH/2
add_rect(s, Inches(fiberX), Inches(fiberY-0.02), Inches(fiberEndX-fiberX), Inches(0.04), fill=ACCENT)
add_textbox(s, Inches(fiberX), Inches(fiberY-0.5), Inches(fiberEndX-fiberX), Inches(0.3), [("隐形光纤  主干",{"size":9,"color":MUTED})], align="center")
# 4个从光猫 串行
rooms=["客厅","主卧","书房","儿童房"]
slaveW=1.5; slaveH=1.15
gap=(fiberEndX-fiberX-4*slaveW)/3
for i,rn in enumerate(rooms):
    rx=fiberX+i*(slaveW+gap)
    ry=fiberY+0.3
    draw_card(s, Inches(rx), Inches(ry), Inches(slaveW), Inches(slaveH), line=RED2, line_w=1.2)
    add_textbox(s, Inches(rx), Inches(ry+0.1), Inches(slaveW), Inches(0.3), [("WiFi6",{"size":10,"bold":True,"color":RED2})], align="center", valign="middle")
    add_textbox(s, Inches(rx), Inches(ry+0.4), Inches(slaveW), Inches(0.35), [("从光猫",{"size":12,"bold":True,"color":BLACK})], align="center", valign="middle")
    add_textbox(s, Inches(rx), Inches(ry+0.75), Inches(slaveW), Inches(0.3), [(rn,{"size":10,"color":MUTED})], align="center", valign="middle")
# 三要素
factors=[("主光猫","WiFi6智能网关\n统一接入与调度",RED2),("隐形光纤","透明/微径光缆\n施工无感知",ACCENT),("从光猫","WiFi6扩展\n按需1拖N部署",BLUE)]
fY=5.95; fH=1.1; fW=(12.43-2*0.3)/3
for i,(t,d,c) in enumerate(factors):
    fx=0.45+i*(fW+0.3)
    draw_card(s, Inches(fx), Inches(fY), Inches(fW), Inches(fH), line=BORDER, line_w=1)
    add_rect(s, Inches(fx), Inches(fY), Inches(0.1), Inches(fH), fill=c)
    add_textbox(s, Inches(fx+0.25), Inches(fY+0.12), Inches(fW-0.4), Inches(0.4), [(t,{"size":13,"bold":True,"color":c})], valign="middle")
    add_textbox(s, Inches(fx+0.25), Inches(fY+0.5), Inches(fW-0.4), Inches(0.55), [(d,{"size":10,"color":MUTED})], valign="top")

# ---------- P5 宽带三次跃迁 ----------
s = new_slide(); hdr(s,"认识FTTR","宽带三次跃迁")
add_textbox(s, SX, SY, SW, Inches(0.4), [("家庭网络的“三次跃迁”：从有网可用，到好用稳定，再到FTTR打通最后十米",{"size":12,"color":MUTED})])
lineY=3.4; lsX=0.45+1.7; leX=0.45+12.43-1.7
add_rect(s, Inches(lsX), Inches(lineY-0.03), Inches(leX-lsX), Inches(0.06), fill=BORDER)
stages=[("1","百兆网关","有网可用","解决上网基础需求\n满足网页浏览、视频观看",MUTED),("2","千兆网关","好用稳定","满足多设备并发\n支撑4K视频、远程办公",BLUE),("3","FTTR全光","最后十米","千兆直达每个房间\n双千兆无盲区覆盖",RED2)]
segW=(leX-lsX)/2
for i,(no,name,tag,desc,col) in enumerate(stages):
    cx=lsX+i*segW
    add_oval(s, Inches(cx-0.32), Inches(lineY-0.32), Inches(0.64), Inches(0.64), fill=col, line=WHITE)
    add_oval(s, Inches(cx-0.32), Inches(lineY-0.32), Inches(0.64), Inches(0.64), fill=col, line=WHITE, line_w=Pt(2) if False else None)
    add_textbox(s, Inches(cx-0.32), Inches(lineY-0.32), Inches(0.64), Inches(0.64), [(no,{"size":22,"bold":True,"color":WHITE,"font":FONT})], align="center", valign="middle")
    add_textbox(s, Inches(cx-1.5), Inches(2.15), Inches(3.0), Inches(0.45), [(name,{"size":16,"bold":True,"color":col})], align="center", valign="middle")
    add_rect(s, Inches(cx-0.65), Inches(2.65), Inches(1.3), Inches(0.38), fill=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_textbox(s, Inches(cx-0.65), Inches(2.65), Inches(1.3), Inches(0.38), [(tag,{"size":11,"bold":True,"color":WHITE})], align="center", valign="middle")
    dY=lineY+0.6; dW=3.0
    draw_card(s, Inches(cx-dW/2), Inches(dY), Inches(dW), Inches(1.6), line=col, line_w=1)
    add_textbox(s, Inches(cx-dW/2+0.2), Inches(dY+0.2), Inches(dW-0.4), Inches(1.2), [(desc,{"size":12,"color":BLACK})], align="center", valign="middle")

# ---------- P6 三大核心优势 ----------
s = new_slide(); hdr(s,"核心卖点","三大核心优势")
add_textbox(s, SX, SY, SW, Inches(0.4), [("覆盖·速率·体验 三大维度层层递进",{"size":12,"color":MUTED})])
cards=[("01","全屋无盲区","COVERAGE",["光纤直达每个房间","WiFi信号无死角覆盖","告别穿墙衰减死角"],"覆盖维度",RED2),
       ("02","千兆到房间","SPEED",["有线无线双千兆","到房速率不打折","实测千兆满格"],"速率维度",BLUE),
       ("03","无缝漫游","EXPERIENCE",["换房间不切路由","视频游戏不掉线","毫秒级无缝切换"],"体验维度",RED2)]
cY=1.85; cH=5.0; cW=(12.43-2*0.3)/3
for i,(no,title,en,pts,foot,col) in enumerate(cards):
    cx=0.45+i*(cW+0.3)
    draw_card(s, Inches(cx), Inches(cY), Inches(cW), Inches(cH), line=BORDER, line_w=1)
    add_rect(s, Inches(cx), Inches(cY), Inches(cW), Inches(1.05), fill=col)
    add_textbox(s, Inches(cx), Inches(cY+0.12), Inches(cW), Inches(0.4), [(en,{"size":20,"bold":True,"color":ACCENT,"font":FONT})], align="center", valign="middle")
    add_textbox(s, Inches(cx), Inches(cY+0.55), Inches(cW), Inches(0.4), [(title,{"size":16,"bold":True,"color":WHITE})], align="center", valign="middle")
    for j,p in enumerate(pts):
        py=cY+1.35+j*0.7
        add_oval(s, Inches(cx+0.3), Inches(py+0.13), Inches(0.14), Inches(0.14), fill=col)
        add_textbox(s, Inches(cx+0.55), Inches(py), Inches(cW-0.75), Inches(0.4), [(p,{"size":12,"color":BLACK})], valign="middle")
    add_rect(s, Inches(cx), Inches(cY+cH-0.5), Inches(cW), Inches(0.5), fill=ICELT)
    add_textbox(s, Inches(cx), Inches(cY+cH-0.5), Inches(cW), Inches(0.5), [(foot,{"size":11,"bold":True,"color":col})], align="center", valign="middle")
add_textbox(s, SX, Inches(6.95), SW, Inches(0.35), [("三大优势层层递进：先有覆盖，再有速率，最后有体验——FTTR价值闭环",{"size":11,"bold":True,"color":RED2})], align="center")

# ---------- P7 传统组网对比 ----------
s = new_slide(); hdr(s,"核心卖点","传统组网对比")
add_textbox(s, SX, SY, SW, Inches(0.4), [("同样的千兆入户，传统组网“最后一十米”打折，FTTR全光全程不打折",{"size":12,"color":MUTED})])
colW=(12.43-0.4)/2; colH=4.5; colY=1.95
# 左：传统
lx=0.45
draw_card(s, Inches(lx), Inches(colY), Inches(colW), Inches(colH), fill=RGBColor(0xFF,0xF5,0xF5), line=RGBColor(0xE8,0xA0,0xA0), line_w=1.2)
add_rect(s, Inches(lx), Inches(colY), Inches(colW), Inches(0.6), fill=RGBColor(0xE8,0xA0,0xA0))
add_textbox(s, Inches(lx+0.25), Inches(colY), Inches(colW-0.5), Inches(0.6), [("✗  传统组网（路由器/电力猫）",{"size":14,"bold":True,"color":WHITE})], valign="middle")
pains=[("穿墙衰减","多道墙后网速大幅衰减，远端房间仅剩几十兆"),("信号死角","卫生间、阳台、二楼等角落无信号覆盖"),("切换掉线","换房间切换路由时视频卡顿、游戏断线"),("速率打折","入户千兆，到房间往往只剩百兆甚至更低")]
for i,(t,d) in enumerate(pains):
    py=colY+0.85+i*0.88
    add_rect(s, Inches(lx+0.25), Inches(py), Inches(0.08), Inches(0.72), fill=RGBColor(0xD0,0x60,0x60))
    add_textbox(s, Inches(lx+0.5), Inches(py), Inches(colW-0.8), Inches(0.38), [(t,{"size":13,"bold":True,"color":RGBColor(0x8B,0x00,0x12)})], valign="middle")
    add_textbox(s, Inches(lx+0.5), Inches(py+0.36), Inches(colW-0.8), Inches(0.36), [(d,{"size":11,"color":RGBColor(0x66,0x66,0x66)})], valign="middle")
# 右：FTTR
rxR=0.45+colW+0.4
draw_card(s, Inches(rxR), Inches(colY), Inches(colW), Inches(colH), fill=RGBColor(0xF0,0xFD,0xF4), line=GREEN, line_w=1.5)
add_rect(s, Inches(rxR), Inches(colY), Inches(colW), Inches(0.6), fill=GREEN)
add_textbox(s, Inches(rxR+0.25), Inches(colY), Inches(colW-0.5), Inches(0.6), [("✓  FTTR 全光组网",{"size":14,"bold":True,"color":WHITE})], valign="middle")
wins=[("光纤到房","光纤直达每个房间，速率无衰减"),("全屋覆盖","双千兆无死角，每个角落信号满格"),("无缝漫游","换房间不切换，视频游戏不掉线"),("到房千兆","入户千兆，到房仍是千兆不打折")]
for i,(t,d) in enumerate(wins):
    py=colY+0.85+i*0.88
    add_rect(s, Inches(rxR+0.25), Inches(py), Inches(0.08), Inches(0.72), fill=GREEN)
    add_textbox(s, Inches(rxR+0.5), Inches(py), Inches(colW-0.8), Inches(0.38), [(t,{"size":13,"bold":True,"color":RGBColor(0x14,0x6C,0x2E)})], valign="middle")
    add_textbox(s, Inches(rxR+0.5), Inches(py+0.36), Inches(colW-0.8), Inches(0.36), [(d,{"size":11,"color":RGBColor(0x44,0x44,0x44)})], valign="middle")
add_rect(s, SX, Inches(6.7), SW, Inches(0.5), fill=RED2, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
add_textbox(s, Inches(0.65), Inches(6.7), Inches(12.0), Inches(0.5),
            [("核心差距：",{"size":13,"bold":True,"color":WHITE}),
             ("传统组网“入户千兆、到房打折”，FTTR“入户千兆、到房仍千兆”——这是卖点的本质",{"size":12,"color":WHITE})], valign="middle")

# ---------- P8 五大应用场景 ----------
s = new_slide(); hdr(s,"应用场景","五大典型场景")
add_textbox(s, SX, SY, SW, Inches(0.4), [("FTTR支撑的高带宽、低时延、多并发场景，传统组网难以胜任",{"size":12,"color":MUTED})])
gridY=1.85; gridH=2.5; gridW=(12.43-0.4)/2; gridGap=0.3; cellH=(gridH-gridGap)/2
scenes=[("4K","4K/8K超清视频","多路超清同时播放，缓冲不卡顿"),("GAME","云游戏低延迟","云端渲染实时交互，时延不掉包"),("AI","智能家居联动","几十设备同时在线，指令秒响应"),("OFFICE","远程办公稳定","视频会议+大文件传输双顺畅")]
for i,(icon,title,desc) in enumerate(scenes):
    col=i%2; row=i//2
    cx=0.45+col*(gridW+0.4); cy=gridY+row*(cellH+gridGap)
    draw_card(s, Inches(cx), Inches(cy), Inches(gridW), Inches(cellH), line=BORDER, line_w=1)
    add_rect(s, Inches(cx), Inches(cy), Inches(1.3), Inches(cellH), fill=RED2)
    add_textbox(s, Inches(cx), Inches(cy), Inches(1.3), Inches(cellH), [(icon,{"size":20,"bold":True,"color":WHITE,"font":FONT})], align="center", valign="middle")
    add_textbox(s, Inches(cx+1.5), Inches(cy+0.2), Inches(gridW-1.7), Inches(0.4), [(title,{"size":15,"bold":True,"color":BLACK})], valign="middle")
    add_textbox(s, Inches(cx+1.5), Inches(cy+0.65), Inches(gridW-1.7), Inches(0.4), [(desc,{"size":11,"color":MUTED})], valign="middle")
eduY=gridY+gridH+0.25; eduH=1.2
draw_card(s, SX, Inches(eduY), SW, Inches(eduH), line=RED2, line_w=1.2, fill=WHITE)
add_rect(s, SX, Inches(eduY), Inches(1.5), Inches(eduH), fill=RED2)
add_textbox(s, SX, Inches(eduY), Inches(1.5), Inches(eduH), [("EDU",{"size":20,"bold":True,"color":WHITE,"font":FONT})], align="center", valign="middle")
add_textbox(s, Inches(0.45+1.7), Inches(eduY+0.15), Inches(12.43-2.0), Inches(0.4), [("在线教育多端并发",{"size":15,"bold":True,"color":RED2})], valign="middle")
add_textbox(s, Inches(0.45+1.7), Inches(eduY+0.6), Inches(12.43-2.0), Inches(0.4), [("孩子上网课、家长视频会议、智能家居联动同时进行，全屋网络不抢不卡",{"size":11,"color":BLACK})], valign="middle")

# ---------- P9 目标客群画像 ----------
s = new_slide(); hdr(s,"销售打法","目标客群画像")
add_textbox(s, SX, SY, SW, Inches(0.4), [("精准识别三类高潜客群，前置触达，提升FTTR转化率",{"size":12,"color":MUTED})])
groups=[("客群一","大户型/复式/别墅",RED2,"多房间多楼层，传统组网覆盖盲区大",["全屋无死角覆盖","楼层间信号贯通","多房间并发不卡"],"高ARPU，价格敏感度低"),
        ("客群二","智能家居重度用户",BLUE,"几十台智能设备同时在线，对并发要求高",["多设备稳定连接","指令低时延响应","设备频繁切换不掉线"],"中高ARPU，看重体验"),
        ("客群三","高带宽需求家庭",RED2,"4K/8K视频、云游戏、远程办公多场景叠加",["千兆到房间不掉速","低时延游戏不掉包","上行带宽充足"],"中高ARPU，愿为体验付费")]
gY=1.85; gH=4.8; gW=(12.43-2*0.3)/3
for i,(tag,title,col,feat,need,price) in enumerate(groups):
    cx=0.45+i*(gW+0.3)
    draw_card(s, Inches(cx), Inches(gY), Inches(gW), Inches(gH), line=BORDER, line_w=1)
    add_rect(s, Inches(cx), Inches(gY), Inches(gW), Inches(0.9), fill=col)
    add_textbox(s, Inches(cx), Inches(gY+0.08), Inches(gW), Inches(0.28), [(tag,{"size":10,"color":WHITE})], align="center", valign="middle")
    add_textbox(s, Inches(cx), Inches(gY+0.38), Inches(gW), Inches(0.5), [(title,{"size":14,"bold":True,"color":WHITE})], align="center", valign="middle")
    add_textbox(s, Inches(cx+0.25), Inches(gY+1.05), Inches(gW-0.5), Inches(0.28), [("客户特征",{"size":10,"bold":True,"color":col})], valign="middle")
    add_textbox(s, Inches(cx+0.25), Inches(gY+1.35), Inches(gW-0.5), Inches(0.6), [(feat,{"size":11,"color":BLACK})], valign="top")
    add_textbox(s, Inches(cx+0.25), Inches(gY+2.05), Inches(gW-0.5), Inches(0.28), [("核心需求",{"size":10,"bold":True,"color":col})], valign="middle")
    for j,n in enumerate(need):
        add_oval(s, Inches(cx+0.28), Inches(gY+2.4+j*0.38+0.08), Inches(0.1), Inches(0.1), fill=col)
        add_textbox(s, Inches(cx+0.48), Inches(gY+2.4+j*0.38), Inches(gW-0.7), Inches(0.32), [(n,{"size":11,"color":BLACK})], valign="middle")
    add_rect(s, Inches(cx), Inches(gY+gH-0.55), Inches(gW), Inches(0.55), fill=ICELT)
    add_textbox(s, Inches(cx), Inches(gY+gH-0.55), Inches(gW), Inches(0.55), [(price,{"size":11,"bold":True,"color":col})], align="center", valign="middle")

# ---------- P10 话术与异议处理 ----------
s = new_slide(); hdr(s,"销售打法","话术与异议处理")
add_textbox(s, SX, SY, SW, Inches(0.4), [("异议处理三步法：算账→体验→服务",{"size":12,"color":MUTED})])
rows=[("“FTTR太贵了”","算账：对比传统组网反复购买扩展器成本，FTTR一次到位、十年免升级"),
      ("“现在网速够用了”","体验：现场测速对比，演示穿墙后到房实际速率"),
      ("“担心施工破坏装修”","服务：隐形光纤透明无感，专业装维全程负责"),
      ("“路由器也能覆盖”","算账：传统组网到房打折，FTTR到房仍千兆，体验完全不同")]
# 表头
tY=1.9; rowH=0.95
add_rect(s, SX, Inches(tY), Inches(4.3), Inches(0.5), fill=RED2); add_textbox(s, Inches(0.65), Inches(tY), Inches(4.0), Inches(0.5), [("客户异议",{"size":13,"bold":True,"color":WHITE})], valign="middle")
add_rect(s, Inches(0.45+4.4), Inches(tY), Inches(8.0), Inches(0.5), fill=BLUE); add_textbox(s, Inches(0.65+4.4), Inches(tY), Inches(7.7), Inches(0.5), [("应对话术与价值传递",{"size":13,"bold":True,"color":WHITE})], valign="middle")
for i,(obq,ans) in enumerate(rows):
    ry=tY+0.5+i*rowH
    col2=RED2 if i%2==0 else BLUE
    add_rect(s, SX, Inches(ry), Inches(4.3), Inches(rowH), fill=WHITE, line=BORDER, line_w=Pt(0.75))
    add_rect(s, SX, Inches(ry), Inches(0.08), Inches(rowH), fill=col2)
    add_textbox(s, Inches(0.65), Inches(ry), Inches(4.0), Inches(rowH), [(obq,{"size":12,"bold":True,"color":col2})], valign="middle")
    add_rect(s, Inches(0.45+4.4), Inches(ry), Inches(8.0), Inches(rowH), fill=WHITE, line=BORDER, line_w=Pt(0.75))
    add_textbox(s, Inches(0.65+4.4), Inches(ry), Inches(7.7), Inches(rowH), [(ans,{"size":11,"color":BLACK})], valign="middle")
add_textbox(s, SX, Inches(6.75), SW, Inches(0.45), [("→ 算账           → 体验           → 服务",{"size":13,"bold":True,"color":RED2})], align="center")

# ---------- P11 主推套餐推荐 ----------
s = new_slide(); hdr(s,"销售打法","主推套餐推荐")
add_textbox(s, SX, SY, SW, Inches(0.4), [("按客群价值分层推荐，主推双千兆档，尊享档锁定高净值客户",{"size":12,"color":MUTED})])
plans=[("入门","千兆起步档",BLUE,"首次迁转FTTR客户\n小户型/中等户型",["千兆宽带","FTTR基础组网","主光猫1台","基础WiFi调优"],"适合尝鲜"),
       ("主推","双千兆主推档",RED2,"大户型/智能家居\n中高ARPU家庭",["千兆宽带","FTTR 1拖2组网","全屋WiFi调优","专属装维服务"],"★ 重点推荐"),
       ("旗舰","全光尊享档",RGBColor(0xA8,0x00,0x1A),"别墅/复式\n高净值家庭",["千兆/超千兆宽带","FTTR 1拖N组网","天翼云盘权益","VIP装维+定期巡检"],"锁定高净值")]
pY=1.9; pH=4.6; pW=(12.43-2*0.3)/3
for i,(lv,name,col,target,rights,tag) in enumerate(plans):
    cx=0.45+i*(pW+0.3)
    draw_card(s, Inches(cx), Inches(pY), Inches(pW), Inches(pH), line=col, line_w=2 if lv=="主推" else 1)
    add_rect(s, Inches(cx), Inches(pY), Inches(pW), Inches(1.0), fill=col)
    add_textbox(s, Inches(cx), Inches(pY+0.12), Inches(pW), Inches(0.28), [(lv,{"size":10,"color":ACCENT})], align="center", valign="middle")
    add_textbox(s, Inches(cx), Inches(pY+0.4), Inches(pW), Inches(0.5), [(name,{"size":15,"bold":True,"color":WHITE})], align="center", valign="middle")
    tagW=1.6
    add_rect(s, Inches(cx+(pW-tagW)/2), Inches(pY+1.12), Inches(tagW), Inches(0.35), fill=ICELT, line=col, line_w=1, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_textbox(s, Inches(cx+(pW-tagW)/2), Inches(pY+1.12), Inches(tagW), Inches(0.35), [(tag,{"size":10,"bold":True,"color":col})], align="center", valign="middle")
    add_textbox(s, Inches(cx+0.25), Inches(pY+1.65), Inches(pW-0.5), Inches(0.25), [("目标客群",{"size":10,"bold":True,"color":col})], valign="middle")
    add_textbox(s, Inches(cx+0.25), Inches(pY+1.95), Inches(pW-0.5), Inches(0.65), [(target,{"size":11,"color":BLACK})], valign="top")
    add_textbox(s, Inches(cx+0.25), Inches(pY+2.75), Inches(pW-0.5), Inches(0.25), [("核心权益",{"size":10,"bold":True,"color":col})], valign="middle")
    for j,r in enumerate(rights):
        add_oval(s, Inches(cx+0.28), Inches(pY+3.18+j*0.33+0.07), Inches(0.1), Inches(0.1), fill=col)
        add_textbox(s, Inches(cx+0.48), Inches(pY+3.18+j*0.33), Inches(pW-0.7), Inches(0.3), [(r,{"size":11,"color":BLACK})], valign="middle")
add_textbox(s, SX, Inches(6.78), SW, Inches(0.3), [("资费方案以当地营业厅公示为准；推荐话术：按户型+设备数+用网习惯三维度匹配套餐档位",{"size":9,"color":MUTED})])

# ---------- P12 要点回顾 ----------
s = new_slide(); hdr(s,"培训回顾","要点总结")
add_textbox(s, SX, SY, SW, Inches(0.4), [("一张图回顾本次培训四大模块核心要点",{"size":12,"color":MUTED})])
reviews=[("01","认识FTTR",RED2,["光纤到房间，全屋双千兆","1拖N组网，主光猫+隐形光纤","家庭宽带第三次跃迁"]),
         ("02","核心卖点",BLUE,["全屋无盲区·千兆到房·无缝漫游","传统组网“到房打折”","FTTR“到房仍千兆”"]),
         ("03","应用场景",RED2,["4K/8K·云游戏·智能家居","远程办公·在线教育","高带宽低时延多并发"]),
         ("04","销售打法",BLUE,["三类客群精准触达","异议处理三步法：算账+体验+服务","入门/主推/尊享三档套餐"])]
rY2=1.65; rH=2.35; twoW=(12.43-0.4)/2
for i,(no,title,col,pts) in enumerate(reviews):
    c2=i%2; r2=i//2
    cx=0.45+c2*(twoW+0.4); cy=rY2+r2*(rH+0.3)
    draw_card(s, Inches(cx), Inches(cy), Inches(twoW), Inches(rH), line=BORDER, line_w=1)
    add_rect(s, Inches(cx), Inches(cy), Inches(1.2), Inches(rH), fill=col)
    add_textbox(s, Inches(cx), Inches(cy+0.3), Inches(1.2), Inches(1.0), [(no,{"size":30,"bold":True,"color":WHITE,"font":FONT})], align="center", valign="middle")
    add_textbox(s, Inches(cx+1.4), Inches(cy+0.3), Inches(twoW-1.6), Inches(0.5), [(title,{"size":17,"bold":True,"color":col})], valign="middle")
    add_rect(s, Inches(cx+1.4), Inches(cy+0.85), Inches(0.8), Inches(0.04), fill=col)
    for j,p in enumerate(pts):
        add_oval(s, Inches(cx+1.43), Inches(cy+1.1+j*0.42+0.08), Inches(0.12), Inches(0.12), fill=col)
        add_textbox(s, Inches(cx+1.65), Inches(cy+1.1+j*0.42), Inches(twoW-1.9), Inches(0.38), [(p,{"size":12,"color":BLACK})], valign="middle")
add_rect(s, SX, Inches(6.7), SW, Inches(0.6), fill=RED2, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
add_textbox(s, Inches(0.65), Inches(6.7), Inches(12.0), Inches(0.6),
            [("学以致用  ·  持续深耕  —  ",{"size":13,"bold":True,"color":ACCENT}),
             ("把FTTR的全屋千兆价值，讲进每一个客户心里",{"size":12,"bold":True,"color":WHITE})], valign="middle")

prs.save(OUT)
print("SAVED:", OUT)
print("slides:", len(prs.slides))
