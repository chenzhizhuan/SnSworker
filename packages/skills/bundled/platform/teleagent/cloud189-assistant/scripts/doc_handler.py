#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
天翼云盘文档小秘书 - 文档处理脚本
支持从 .docx / .txt / .md / .pdf 提取文本，以及将文本回写为 .docx / .txt / .md
"""

import argparse
import json
import os
import sys


def extract_from_docx(file_path):
    """从 .docx 文件提取文本，保留段落结构"""
    from docx import Document
    doc = Document(file_path)
    paragraphs = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            paragraphs.append(text)
    return "\n\n".join(paragraphs)


def extract_from_txt(file_path):
    """从 .txt / .md 文件提取文本"""
    with open(file_path, "r", encoding="utf-8") as f:
        return f.read()


def extract_from_pdf(file_path):
    """从 .pdf 文件提取文本"""
    try:
        from pypdf import PdfReader
    except ImportError:
        try:
            from PyPDF2 import PdfReader
        except ImportError:
            return None, "需要安装 pypdf: pip install pypdf"
    reader = PdfReader(file_path)
    pages = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            pages.append(text.strip())
    return "\n\n".join(pages)


def extract_text(file_path):
    """根据文件扩展名提取文本内容"""
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".docx":
        try:
            return extract_from_docx(file_path), None
        except ImportError:
            return None, "需要安装 python-docx: pip install python-docx"
        except Exception as e:
            return None, str(e)
    elif ext in (".txt", ".md"):
        return extract_from_txt(file_path), None
    elif ext == ".pdf":
        return extract_from_pdf(file_path)
    else:
        return None, f"不支持的文件格式: {ext}（支持 .docx/.txt/.md/.pdf）"


def save_to_docx(content, output_path):
    """将文本保存为 .docx 文件，每个段落作为独立段落"""
    from docx import Document
    doc = Document()
    paragraphs = content.split("\n\n")
    for para in paragraphs:
        para = para.strip()
        if para:
            # 处理段落内的单个换行
            lines = para.split("\n")
            for i, line in enumerate(lines):
                if i == 0:
                    p = doc.add_paragraph(line)
                else:
                    p.add_run("\n" + line)
    doc.save(output_path)


def save_to_txt(content, output_path):
    """将文本保存为 .txt / .md 文件"""
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(content)


def save_document(content, output_path, source_format=None):
    """根据输出路径扩展名或源格式保存文档"""
    ext = os.path.splitext(output_path)[1].lower()
    if ext == ".docx":
        try:
            save_to_docx(content, output_path)
            return None
        except ImportError:
            return "需要安装 python-docx: pip install python-docx"
        except Exception as e:
            return str(e)
    elif ext in (".txt", ".md"):
        save_to_txt(content, output_path)
        return None
    elif source_format == ".pdf":
        # PDF 无法直接回写，保存为 .docx
        docx_path = os.path.splitext(output_path)[0] + ".docx"
        try:
            save_to_docx(content, docx_path)
            return None
        except Exception as e:
            return str(e)
    else:
        # 默认保存为 .docx
        if ext == "":
            output_path += ".docx"
        try:
            save_to_docx(content, output_path)
            return None
        except Exception as e:
            return str(e)


def main():
    parser = argparse.ArgumentParser(
        description="天翼云盘文档小秘书 - 文档文本提取与回写"
    )
    subparsers = parser.add_subparsers(dest="command", help="子命令")

    # extract 子命令
    extract_parser = subparsers.add_parser("extract", help="从文档提取文本")
    extract_parser.add_argument("--input", required=True, help="输入文件路径")

    # save 子命令
    save_parser = subparsers.add_parser("save", help="将文本保存为文档")
    save_parser.add_argument("--content-file", required=True,
                            help="包含处理结果的文本文件路径")
    save_parser.add_argument("--output", required=True, help="输出文件路径")
    save_parser.add_argument("--source-format", default=None,
                            help="源文件格式（如 .pdf），用于决定回写策略")

    args = parser.parse_args()

    if args.command == "extract":
        if not os.path.exists(args.input):
            print(json.dumps({"success": False, "message": f"文件不存在: {args.input}"},
                             ensure_ascii=False))
            sys.exit(1)
        text, error = extract_text(args.input)
        if error:
            print(json.dumps({"success": False, "message": error},
                             ensure_ascii=False))
            sys.exit(1)
        print(json.dumps({
            "success": True,
            "file": args.input,
            "format": os.path.splitext(args.input)[1].lower(),
            "char_count": len(text),
            "content": text
        }, ensure_ascii=False))

    elif args.command == "save":
        if not os.path.exists(args.content_file):
            print(json.dumps({"success": False,
                              "message": f"内容文件不存在: {args.content_file}"},
                             ensure_ascii=False))
            sys.exit(1)
        with open(args.content_file, "r", encoding="utf-8") as f:
            content = f.read()
        error = save_document(content, args.output, args.source_format)
        if error:
            print(json.dumps({"success": False, "message": error},
                             ensure_ascii=False))
            sys.exit(1)
        actual_path = args.output
        if args.source_format == ".pdf" and os.path.splitext(args.output)[1].lower() != ".docx":
            actual_path = os.path.splitext(args.output)[0] + ".docx"
        print(json.dumps({
            "success": True,
            "message": "文档保存成功",
            "output_path": os.path.abspath(actual_path),
            "file_size": os.path.getsize(actual_path)
        }, ensure_ascii=False))

    else:
        parser.print_help()


if __name__ == "__main__":
    main()
