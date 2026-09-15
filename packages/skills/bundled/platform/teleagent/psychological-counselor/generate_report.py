#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
心理疏导师 - Word报告生成器
生成完整的心理疏导记录与成长报告
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from datetime import datetime

class PsychologicalReport:
    def __init__(self, session_data):
        self.doc = Document()
        self.data = session_data
        self.setup_styles()
    
    def setup_styles(self):
        """设置文档样式"""
        self.doc.styles['Normal'].font.name = 'Microsoft YaHei'
        self.doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        self.doc.styles['Normal'].font.size = Pt(11)
    
    def add_title(self, text, level=1):
        """添加标题"""
        p = self.doc.add_paragraph()
        if level == 1:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text)
            run.font.size = Pt(20)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 51, 102)
        elif level == 2:
            run = p.add_run(text)
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 102, 153)
        else:
            run = p.add_run(text)
            run.font.size = Pt(12)
            run.font.bold = True
        
        run.font.name = 'Microsoft YaHei'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        return p
    
    def add_paragraph(self, text, bold=False, italic=False):
        """添加段落"""
        p = self.doc.add_paragraph()
        run = p.add_run(text)
        run.font.name = 'Microsoft YaHei'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        run.font.size = Pt(11)
        run.bold = bold
        run.italic = italic
        return p
    
    def add_quote(self, text):
        """添加引用"""
        self.doc.add_paragraph()
        p = self.doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.right_indent = Inches(0.3)
        run = p.add_run(f'"{text}"')
        run.font.name = 'Microsoft YaHei'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        run.font.size = Pt(11)
        run.italic = True
        run.font.color.rgb = RGBColor(80, 80, 80)
        self.doc.add_paragraph()
    
    def add_highlight(self, text):
        """添加高亮"""
        p = self.doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        run = p.add_run(text)
        run.font.name = 'Microsoft YaHei'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        run.font.size = Pt(11)
        run.bold = True
        run.font.color.rgb = RGBColor(0, 102, 153)
        return p
    
    def add_table(self, headers, rows):
        """添加表格"""
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.style = 'Light Grid Accent 1'
        
        # 表头
        hdr_cells = table.rows[0].cells
        for i, header in enumerate(headers):
            hdr_cells[i].text = header
            for paragraph in hdr_cells[i].paragraphs:
                for run in paragraph.runs:
                    run.font.bold = True
                    run.font.name = 'Microsoft YaHei'
                    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        
        # 数据行
        for row_data in rows:
            row_cells = table.add_row().cells
            for i, cell_data in enumerate(row_data):
                row_cells[i].text = str(cell_data)
                for paragraph in row_cells[i].paragraphs:
                    for run in paragraph.runs:
                        run.font.name = 'Microsoft YaHei'
                        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        
        return table
    
    def generate(self):
        """生成完整报告"""
        # 封面
        self.add_title("心理疏导记录与成长报告", level=1)
        
        user_name = self.data.get("user_name", "匿名")
        self.add_title(f"记录者：{user_name}", level=2)
        self.doc.add_paragraph()
        
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"生成日期：{datetime.now().strftime('%Y年%m月%d日')}")
        run.font.name = 'Microsoft YaHei'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(150, 150, 150)
        
        self.doc.add_page_break()
        
        # 引言
        self.add_quote("理解自己，善待自己，你有能力应对一切心理挑战。")
        self.add_paragraph('本报告记录了一次完整的心理疏导过程，遵循"感知-拆解-疏导-修复-巩固"五步法，帮助你理清心理困扰，回归内心平衡。')
        
        # 第一步：感知
        self.add_title("第一步：心理状态感知与界定", level=2)
        
        perception = self.data.get("perception", {})
        
        self.add_title("1.1 心理主体", level=3)
        self.add_paragraph(f"主体类型：{perception.get('subject', '未记录')}")
        
        self.add_title("1.2 心理信号捕捉", level=3)
        signals = perception.get("signals", {})
        self.add_highlight("情绪信号：")
        self.add_paragraph(signals.get("emotion", "未记录"))
        self.add_highlight("行为信号：")
        self.add_paragraph(signals.get("behavior", "未记录"))
        self.add_highlight("认知信号：")
        self.add_paragraph(signals.get("cognition", "未记录"))
        
        self.add_title("1.3 问题属性界定", level=3)
        self.add_paragraph(f"问题类型：{perception.get('problem_type', '未记录')}")
        self.add_paragraph(f"严重程度（1-10）：{perception.get('severity', '未记录')}")
        
        self.doc.add_page_break()
        
        # 第二步：拆解
        self.add_title("第二步：心理问题根源拆解", level=2)
        
        analysis = self.data.get("analysis", {})
        
        self.add_title("2.1 触发事件", level=3)
        self.add_paragraph(analysis.get("trigger", "未记录"))
        
        self.add_title("2.2 深层心理需求", level=3)
        self.add_paragraph(f"核心需求：{analysis.get('need', '未记录')}")
        self.add_paragraph(f"需求强度（1-10）：{analysis.get('need_intensity', '未记录')}")
        
        self.add_title("2.3 认知偏差", level=3)
        self.add_highlight("负面想法：")
        self.add_paragraph(analysis.get("negative_thought", "未记录"))
        self.add_highlight("客观想法：")
        self.add_paragraph(analysis.get("objective_thought", "未记录"))
        
        self.add_title("2.4 可控与不可控", level=3)
        self.add_highlight("可控因素：")
        self.add_paragraph(analysis.get("controllable", "未记录"))
        self.add_highlight("不可控因素（需接纳）：")
        self.add_paragraph(analysis.get("uncontrollable", "未记录"))
        
        self.doc.add_page_break()
        
        # 第三步：疏导
        self.add_title("第三步：针对性心理疏导与情绪释放", level=2)
        
        guidance = self.data.get("guidance", {})
        
        self.add_title("3.1 情绪接纳", level=3)
        self.add_quote(guidance.get("acceptance", "我接纳当下的所有情绪，这是正常的反应。"))
        
        self.add_title("3.2 情绪释放方式", level=3)
        self.add_paragraph(f"采用方式：{guidance.get('release_method', '未记录')}")
        
        self.add_title("3.3 认知纠正", level=3)
        self.add_highlight("原认知偏差：")
        self.add_paragraph(guidance.get("original_cognition", "未记录"))
        self.add_highlight("纠正后认知：")
        self.add_paragraph(guidance.get("corrected_cognition", "未记录"))
        
        self.add_title("3.4 执念放下", level=3)
        self.add_paragraph(f"放下执念：{guidance.get("letting_go", '未记录')}")
        
        self.doc.add_page_break()
        
        # 第四步：修复
        self.add_title("第四步：心理平衡修复与问题应对", level=2)
        
        repair = self.data.get("repair", {})
        
        self.add_title("4.1 理性回归", level=3)
        self.add_paragraph(f"情绪强度变化：{repair.get('emotion_change', '未记录')}（从__分降至__分）")
        
        self.add_title("4.2 应对行动", level=3)
        self.add_highlight("行动一（今天/现在）：")
        self.add_paragraph(repair.get("action_1", "未记录"))
        self.add_highlight("行动二（本周）：")
        self.add_paragraph(repair.get("action_2", "未记录"))
        
        self.add_title("4.3 自我支撑强化", level=3)
        self.add_paragraph("我的优势：")
        strengths = repair.get("strengths", [])
        for i, strength in enumerate(strengths, 1):
            self.add_paragraph(f"{i}. {strength}")
        
        self.add_title("4.4 场景化解决方案", level=3)
        self.add_paragraph(f"适用场景：{repair.get('scene', '未记录')}")
        self.add_paragraph(f"应对方式：{repair.get('solution', '未记录')}")
        
        self.doc.add_page_break()
        
        # 第五步：巩固
        self.add_title("第五步：心理状态巩固与预防", level=2)
        
        consolidation = self.data.get("consolidation", {})
        
        self.add_title("5.1 过程复盘", level=3)
        self.add_highlight("本次收获：")
        self.add_paragraph(consolidation.get("gain", "未记录"))
        self.add_highlight("下次改进：")
        self.add_paragraph(consolidation.get("improvement", "未记录"))
        
        self.add_title("5.2 预警机制", level=3)
        self.add_paragraph(f"我的预警信号：{consolidation.get('warning_signs', '未记录')}")
        
        self.add_title("5.3 韧性提升计划", level=3)
        self.add_paragraph("日常习惯：")
        habits = consolidation.get("habits", [])
        for habit in habits:
            self.add_paragraph(f"• {habit}")
        
        self.add_title("5.4 长期维护", level=3)
        self.add_paragraph(f"维护计划：{consolidation.get('maintenance', '未记录')}")
        
        # 总结
        self.doc.add_page_break()
        self.add_title("疏导总结", level=2)
        
        self.add_highlight("核心问题：")
        self.add_paragraph(self.data.get("summary_problem", "未记录"))
        
        self.add_highlight("关键突破：")
        self.add_paragraph(self.data.get("summary_breakthrough", "未记录"))
        
        self.add_highlight("成长收获：")
        self.add_paragraph(self.data.get("summary_growth", "未记录"))
        
        # 结语
        self.doc.add_paragraph()
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("—" * 20)
        run.font.name = 'Microsoft YaHei'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        run.font.color.rgb = RGBColor(200, 200, 200)
        
        self.doc.add_paragraph()
        self.add_quote("每一次心理疏导，都是一次自我成长的旅程。你已经做得很好，继续前进。")
        
        # 保存
        output_path = f"心理疏导报告_{user_name}_{datetime.now().strftime('%Y%m%d')}.docx"
        self.doc.save(output_path)
        return output_path


def main():
    """主函数 - 示例数据"""
    # 示例数据
    session_data = {
        "user_name": "示例用户",
        "perception": {
            "subject": "自我干预",
            "signals": {
                "emotion": "焦虑（强度8/10），担心工作表现",
                "behavior": "失眠、拖延、回避社交",
                "cognition": "我不行、一定会搞砸、别人都在看我笑话"
            },
            "problem_type": "持续心理困扰（2周）",
            "severity": "8"
        },
        "analysis": {
            "trigger": "即将进行的重要项目汇报",
            "need": "认同感、掌控感",
            "need_intensity": "9",
            "negative_thought": "我必须完美，否则就是失败",
            "objective_thought": "尽力就好，完美不是唯一标准，一次表现不定义我的能力",
            "controllable": "准备程度、演讲练习、心态调整",
            "uncontrollable": "他人评价、最终结果、不可预见的意外"
        },
        "guidance": {
            "acceptance": "我感到焦虑是正常的，这说明我重视这次汇报。我允许自己有这样的情绪。",
            "release_method": "深呼吸调节 + 倾诉表达（与朋友聊聊）",
            "original_cognition": "我必须完美，否则就是失败（绝对化思维）",
            "corrected_cognition": "我可以尽力，允许不完美。完成比完美重要，一次表现不定义我的全部能力。",
            "letting_go": "放下对完美表现的执念，接受足够好的标准"
        },
        "repair": {
            "emotion_change": "从8分降至4分",
            "action_1": "今天：完成汇报PPT的初稿，不求完美",
            "action_2": "本周：每天练习演讲10分钟，找朋友模拟一次",
            "strengths": ["我有丰富的项目经验", "我善于逻辑思维", "我之前成功完成过类似汇报"],
            "scene": "职场压力",
            "solution": "充分准备 + 调整心态 + 允许不完美"
        },
        "consolidation": {
            "gain": "认识到完美主义是焦虑的根源，学会了足够好的标准",
            "improvement": "提前准备，不要等到最后一刻；建立准备-练习-放松的固定流程",
            "warning_signs": "失眠、反复想一件事、自我否定话语",
            "habits": [
                "每天5分钟情绪觉察",
                "每周1次心理复盘",
                "建立可控圈思维习惯",
                "定期运动释放压力"
            ],
            "maintenance": "每月进行一次深度自我对话，每季度回顾成长记录"
        },
        "summary_problem": "对重要汇报的过度焦虑，源于完美主义和对他人评价的过度在意",
        "summary_breakthrough": "认识到完美不是必须，建立了足够好的标准，制定了具体的准备计划",
        "summary_growth": "学会了区分可控与不可控，掌握了认知纠偏的方法，建立了应对职场压力的固定流程"
    }
    
    generator = PsychologicalReport(session_data)
    output_path = generator.generate()
    print(f"报告已生成：{output_path}")


if __name__ == "__main__":
    main()
